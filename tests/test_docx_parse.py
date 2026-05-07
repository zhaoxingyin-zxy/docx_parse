import json
import re
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

from docx_parse.backend.office.office_middle_json_mkcontent import union_make
from docx_parse.model.docx.docx_converter import DocxConverter
from docx_parse.utils.enum_class import MakeMode
from docx_parse import (
    convert_docx_file_to_jsonl,
    convert_docx_file_to_markdown,
    convert_docx_file_to_middle_json,
    convert_docx_file_to_model_output,
)


DOCX_01_PATH = Path(r"D:\zxy\code\MinerU-master\MinerU-master\demo\office_docs\docx_01.docx")
DOCX_01_TABLE_BASELINE = Path(__file__).parent / "fixtures" / "docx_01_tables_baseline.json"


def _table_records_from_jsonl(jsonl: str) -> list[dict]:
    return [
        record
        for record in (json.loads(line) for line in jsonl.splitlines())
        if record.get("type") == "table"
    ]


def _table_signature(html: str) -> dict:
    soup = BeautifulSoup(html, "html.parser")
    cells = soup.find_all(["td", "th"])
    text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True)).strip()
    text = re.sub(r"\s+([,.;:!?，。；：！？])", r"\1", text)
    return {
        "tables": len(soup.find_all("table")),
        "rows": len(soup.find_all("tr")),
        "cells": len(cells),
        "images": len(soup.find_all("img")),
        "text": text,
        "colspans": [cell.get("colspan") for cell in cells if cell.get("colspan")],
        "rowspans": [cell.get("rowspan") for cell in cells if cell.get("rowspan")],
    }


def _make_sample_docx(path: Path, image_path: Path) -> None:
    image = Image.new("RGB", (96, 48), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((8, 8, 88, 40), outline="black", width=2)
    draw.text((20, 17), "DOCX", fill="black")
    image.save(image_path)

    doc = Document()
    doc.add_heading("Standalone DOCX Parser", level=1)

    paragraph = doc.add_paragraph("This parser keeps ")
    bold = paragraph.add_run("bold")
    bold.bold = True
    paragraph.add_run(" and ")
    italic = paragraph.add_run("italic")
    italic.italic = True
    paragraph.add_run(" text.")

    doc.add_heading("Data", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Answer"
    table.cell(1, 1).text = "42"

    doc.add_picture(str(image_path))
    doc.add_paragraph("Figure 1: generated test image", style="Caption")
    doc.save(path)


def _set_outline_level(paragraph, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def _set_table_row_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)


def _append_simple_omml_text(paragraph, text: str) -> None:
    o_math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    o_math.append(math_run)
    paragraph._p.append(o_math)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _append_sdt_text(paragraph, text: str) -> None:
    sdt = OxmlElement("w:sdt")
    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    sdt_content.append(run)
    sdt.append(sdt_content)
    paragraph._p.append(sdt)


def test_docx_to_model_output_middle_json_markdown_and_jsonl(tmp_path):
    docx_path = tmp_path / "sample.docx"
    source_image = tmp_path / "source.png"
    image_output_dir = tmp_path / "images"
    _make_sample_docx(docx_path, source_image)

    model_output = convert_docx_file_to_model_output(docx_path)
    assert isinstance(model_output, list)
    assert any(block.get("type") == "title" for page in model_output for block in page)
    assert any(block.get("type") == "table" for page in model_output for block in page)
    assert any(block.get("type") == "image" for page in model_output for block in page)

    middle_json = convert_docx_file_to_middle_json(docx_path, image_output_dir=image_output_dir)
    assert middle_json["_backend"] == "office"
    assert middle_json["pdf_info"]
    assert any(image_output_dir.iterdir())

    markdown = convert_docx_file_to_markdown(
        docx_path,
        image_output_dir=image_output_dir,
        image_dir_name="images",
    )
    assert "# " in markdown
    assert "Standalone DOCX Parser" in markdown
    assert "## " in markdown
    assert "Data" in markdown
    assert "<table>" in markdown
    assert "Answer" in markdown
    assert "42" in markdown
    assert "![](images/" in markdown
    assert "Figure 1: generated test image" in markdown
    assert "**bold**" in markdown
    assert "*italic*" in markdown

    jsonl = convert_docx_file_to_jsonl(
        docx_path,
        image_output_dir=image_output_dir,
        image_dir_name="images",
    )
    records = [json.loads(line) for line in jsonl.splitlines()]
    assert records
    assert all(isinstance(record, dict) for record in records)
    assert any(record.get("text") and "Standalone DOCX Parser" in record["text"] for record in records)
    assert any(record.get("type") == "table" and "Answer" in record.get("table_body", "") for record in records)
    assert any(record.get("type") == "image" and record.get("img_path", "").startswith("images/") for record in records)
    assert all("page_idx" in record for record in records)


def test_outline_level_9_is_treated_as_body_text(tmp_path):
    docx_path = tmp_path / "outline_body.docx"
    doc = Document()
    heading = doc.add_heading("Real Heading", level=2)
    assert heading.text == "Real Heading"
    paragraph = doc.add_paragraph("Body paragraph with DOCX outline level 9.")
    _set_outline_level(paragraph, 9)
    doc.save(docx_path)

    markdown = convert_docx_file_to_markdown(docx_path)
    jsonl = convert_docx_file_to_jsonl(docx_path)
    records = [json.loads(line) for line in jsonl.splitlines()]

    assert "## " in markdown
    assert "Real Heading" in markdown
    assert "##########" not in markdown
    assert "Body paragraph with DOCX outline level 9." in markdown
    assert not any(
        record.get("text") == "Body paragraph with DOCX outline level 9."
        and record.get("text_level") == 10
        for record in records
    )


def test_table_cell_line_breaks_and_empty_paragraphs_are_preserved(tmp_path):
    docx_path = tmp_path / "table_breaks.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    first = cell.paragraphs[0]
    first.add_run("first line")
    first.add_run().add_break()
    first.add_run("second line")
    cell.add_paragraph("")
    cell.add_paragraph("after empty paragraph")
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    table_record = next(record for record in records if record.get("type") == "table")
    table_body = table_record["table_body"]

    assert "<p>" not in table_body
    assert "first line<br/>second line<br/><br/>after empty paragraph" in table_body
    assert "after empty paragraph<br/></th>" in table_body


def test_table_header_cells_are_rendered_as_th(tmp_path):
    docx_path = tmp_path / "table_header.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    _set_table_row_header(table.rows[0])
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Answer"
    table.cell(1, 1).text = "42"
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    table_record = next(record for record in records if record.get("type") == "table")
    table_body = table_record["table_body"]

    assert "<thead><tr><th>Name<br/></th><th>Value<br/></th></tr></thead>" in table_body
    assert "<tbody>" not in table_body
    assert "<th>Name<br/></th>" in table_body
    assert "<th>Value<br/></th>" in table_body
    assert "<td>Answer<br/></td>" in table_body
    assert "<td>42<br/></td>" in table_body


def test_empty_table_cell_paragraph_is_preserved_as_break(tmp_path):
    docx_path = tmp_path / "empty_table_cell.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = ""
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    table_record = next(record for record in records if record.get("type") == "table")
    table_body = table_record["table_body"]

    assert "<p>" not in table_body
    assert "<td><br/></td>" in table_body


def test_table_formula_plain_text_is_preserved_without_latex_conversion(tmp_path):
    docx_path = tmp_path / "table_formula_text.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.add_run("formula: ")
    _append_simple_omml_text(paragraph, "A=πr²")
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    table_record = next(record for record in records if record.get("type") == "table")

    assert "formula: A=πr²" in table_record["table_body"]
    assert "<eq>" not in table_record["table_body"]


def test_header_footer_parts_are_preserved(tmp_path):
    docx_path = tmp_path / "header_footer.docx"
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Document Header"
    section.footer.paragraphs[0].text = "Document Footer"
    doc.add_paragraph("Body text")
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]

    assert any(record.get("type") == "header" and record.get("text") == "Document Header" for record in records)
    assert any(record.get("type") == "footer" and record.get("text") == "Document Footer" for record in records)
    assert any(record.get("type") == "text" and record.get("text") == "Body text" for record in records)


def test_fast_path_preserves_hyperlink_and_visible_styles(tmp_path):
    docx_path = tmp_path / "complex_inline.docx"
    doc = Document()
    link_paragraph = doc.add_paragraph("Go to ")
    _add_hyperlink(link_paragraph, "MinerU", "https://github.com/opendatalab/MinerU")

    style_paragraph = doc.add_paragraph()
    underline = style_paragraph.add_run("under")
    underline.underline = True
    style_paragraph.add_run(" ")
    strike = style_paragraph.add_run("strike")
    strike.font.strike = True

    script_paragraph = doc.add_paragraph("x")
    sup = script_paragraph.add_run("2")
    sup.font.superscript = True
    script_paragraph.add_run(" H")
    sub = script_paragraph.add_run("2")
    sub.font.subscript = True
    script_paragraph.add_run("O")
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    texts = [record.get("text", "") for record in records if record.get("type") == "text"]
    joined = "\n".join(texts)

    assert "[MinerU](https://github.com/opendatalab/MinerU)" in joined
    assert "<u>under</u>" in joined
    assert "~~strike~~" in joined
    assert "x2 H2O" in joined


def test_fast_path_preserves_sdt_and_body_formula_text(tmp_path):
    docx_path = tmp_path / "sdt_formula.docx"
    doc = Document()
    sdt_paragraph = doc.add_paragraph("prefix ")
    _append_sdt_text(sdt_paragraph, "controlled text")
    formula_paragraph = doc.add_paragraph("formula ")
    _append_simple_omml_text(formula_paragraph, "A=πr²")
    doc.save(docx_path)

    records = [json.loads(line) for line in convert_docx_file_to_jsonl(docx_path).splitlines()]
    texts = [record.get("text", "") for record in records if record.get("type") == "text"]
    joined = "\n".join(texts)

    assert "prefix controlled text" in joined
    assert "formula " in joined
    assert any("A=" in text for text in texts)


def test_tolerant_mode_skips_failed_paragraph_and_records_error(tmp_path, monkeypatch):
    docx_path = tmp_path / "tolerant_paragraph.docx"
    doc = Document()
    doc.add_paragraph("before")
    doc.add_paragraph("bad paragraph")
    doc.add_paragraph("after")
    doc.save(docx_path)

    original_handle_text_elements = DocxConverter._handle_text_elements

    def fail_bad_paragraph(self, element):
        paragraph = DocumentParagraph(element, self.docx_obj)
        if paragraph.text == "bad paragraph":
            raise RuntimeError("synthetic paragraph failure")
        return original_handle_text_elements(self, element)

    from docx.text.paragraph import Paragraph as DocumentParagraph

    monkeypatch.setattr(DocxConverter, "_handle_text_elements", fail_bad_paragraph)

    with pytest.raises(RuntimeError, match="synthetic paragraph failure"):
        convert_docx_file_to_jsonl(docx_path, tolerant=False)

    middle_json = convert_docx_file_to_middle_json(docx_path)
    assert middle_json.get("_parse_errors")
    assert middle_json["_parse_errors"][0]["stage"] == "paragraph"

    jsonl = convert_docx_file_to_jsonl(docx_path)
    texts = [
        record.get("text", "")
        for record in (json.loads(line) for line in jsonl.splitlines())
        if record.get("type") == "text"
    ]
    assert "before" in texts
    assert "after" in texts
    assert "bad paragraph" not in texts


def test_tolerant_union_make_skips_malformed_content_block():
    pdf_info = [
        {
            "page_idx": 0,
            "para_blocks": [
                {"type": "text", "lines": [{"spans": [{"type": "text", "content": "ok"}]}]},
                {"type": "image", "blocks": "not-a-list"},
            ],
            "discarded_blocks": [],
        }
    ]

    with pytest.raises(TypeError):
        union_make(pdf_info, MakeMode.CONTENT_LIST, tolerant=False)

    records = union_make(pdf_info, MakeMode.CONTENT_LIST)
    assert records == [{"type": "text", "text": "ok", "page_idx": 0}]


def _legacy_formula_text_parts(text: str) -> list[str]:
    return [
        re.sub(r"\s+", " ", part).strip()
        for part in re.split(r"\$[^$]*\$", text)
        if part.strip()
    ]


def test_docx_01_table_output_matches_legacy_baseline():
    if not DOCX_01_PATH.exists():
        pytest.skip(f"local regression DOCX is missing: {DOCX_01_PATH}")

    baseline_tables = json.loads(DOCX_01_TABLE_BASELINE.read_text(encoding="utf-8"))
    current_tables = _table_records_from_jsonl(convert_docx_file_to_jsonl(DOCX_01_PATH))

    assert len(current_tables) == len(baseline_tables) == 8
    for current, baseline in zip(current_tables, baseline_tables):
        assert current.get("table_caption", []) == baseline.get("table_caption", [])
        assert current.get("page_idx") == baseline.get("page_idx")
        current_signature = _table_signature(current["table_body"])
        baseline_signature = _table_signature(baseline["table_body"])
        for key in ["tables", "rows", "cells", "images", "colspans", "rowspans"]:
            assert current_signature[key] == baseline_signature[key]
        if "$" in baseline_signature["text"]:
            for part in _legacy_formula_text_parts(baseline_signature["text"]):
                assert part in current_signature["text"]
        else:
            assert current_signature["text"] == baseline_signature["text"]
    first_table = current_tables[0]["table_body"]
    assert "<strong" not in first_table
    assert "<em" not in first_table
    assert "This is a list:<br/><ul>" in first_table
    assert "This is a formatted list:<br/><ul>" in first_table
    assert "Third paragraph before a numbered list<br/><ol>" in first_table
